"""Unit tests for the unified PDF, DOCX, TXT, and Markdown processor."""

from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import document_processor as processor
import vector_store
from pdf_processor import SCANNED_PDF_ERROR, process_pdf


TEST_VISITOR_ID = "11111111-1111-4111-8111-111111111111"


def make_pdf_bytes(text: str | None = None) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )
        stream = DecodedStreamObject()
        safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream.set_data(
            f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode("latin-1")
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def make_docx_bytes(*, include_content: bool = True) -> bytes:
    document = Document()
    if include_content:
        document.add_heading("交通安排", level=1)
        document.add_paragraph(
            "这是 Word 文档的普通段落，用于验证标题、段落和文本切分功能。" * 3
        )
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "时间"
        table.cell(0, 1).text = "安排"
        table.cell(1, 0).text = "09:00"
        table.cell(1, 1).text = "前往博物馆，并在附近用餐。"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class DocumentProcessorTests(unittest.TestCase):
    def test_pdf_preserves_real_page_locator(self):
        text = "A normal PDF document contains selectable text for retrieval. " * 3
        result = processor.process_document(make_pdf_bytes(text), "guide.pdf")

        self.assertEqual(result["source_type"], "pdf")
        self.assertEqual(result["total_pages"], 1)
        self.assertTrue(result["chunks"])
        chunk = result["chunks"][0]
        self.assertEqual(chunk["page_number"], 1)
        self.assertEqual(chunk["locator_type"], "page")
        self.assertEqual(chunk["locator_value"], "1")

    def test_empty_pdf_keeps_scanned_pdf_error(self):
        with self.assertRaisesRegex(
            processor.DocumentProcessingError,
            SCANNED_PDF_ERROR,
        ):
            processor.process_document(make_pdf_bytes(), "empty.pdf")

    def test_process_pdf_compatibility_entrypoint_has_new_metadata(self):
        text = "Legacy process_pdf callers continue to receive valid page chunks. " * 3
        result = process_pdf(make_pdf_bytes(text), source_file="legacy.pdf")

        self.assertEqual(result["source_type"], "pdf")
        self.assertEqual(result["chunks"][0]["source_type"], "pdf")
        self.assertEqual(result["chunks"][0]["locator_type"], "page")

    def test_docx_reads_heading_paragraph_and_table_in_order(self):
        result = processor.process_document(
            make_docx_bytes(),
            "travel.docx",
        )

        self.assertEqual(result["source_type"], "docx")
        combined = "\n".join(chunk["chunk_text"] for chunk in result["chunks"])
        self.assertIn("交通安排", combined)
        self.assertIn("Word 文档的普通段落", combined)
        self.assertIn("时间 | 安排", combined)
        self.assertIn("09:00 | 前往博物馆", combined)
        self.assertTrue(all(chunk["page_number"] is None for chunk in result["chunks"]))
        self.assertEqual(result["chunks"][0]["locator_type"], "section")
        self.assertEqual(result["chunks"][0]["locator_value"], "交通安排")

    def test_empty_docx_is_rejected(self):
        with self.assertRaisesRegex(
            processor.DocumentProcessingError,
            "没有可用的文本",
        ):
            processor.process_document(make_docx_bytes(include_content=False), "empty.docx")

    def test_txt_supports_utf8_and_bom(self):
        text = "TXT 文档包含用于知识库检索的 UTF-8 文本内容。" * 4
        plain = processor.process_document(text.encode("utf-8"), "notes.txt")
        bom = processor.process_document(
            b"\xef\xbb\xbf" + text.encode("utf-8"),
            "notes-bom.txt",
        )

        for result in (plain, bom):
            self.assertEqual(result["source_type"], "txt")
            self.assertTrue(result["chunks"])
            self.assertTrue(all(chunk["page_number"] is None for chunk in result["chunks"]))
            self.assertTrue(all(chunk["locator_type"] == "chunk" for chunk in result["chunks"]))

    def test_txt_rejects_non_utf8(self):
        with self.assertRaisesRegex(
            processor.DocumentProcessingError,
            "UTF-8",
        ):
            processor.process_document(b"\xff\xfe\xfa", "legacy.txt")

    def test_markdown_preserves_sections_and_removes_noise(self):
        markdown = (
            "# 交通安排\n\n"
            "**地铁** 是主要交通方式，可以查看 [线路图](https://example.invalid)。" * 3
            + "\n\n## 餐饮建议\n\n"
            + "- 选择附近餐厅，并提前确认营业时间。" * 4
            + "\n<script>dangerous()</script>"
        )
        result = processor.process_document(markdown.encode("utf-8"), "plan.md")

        self.assertEqual(result["source_type"], "md")
        locators = {chunk["locator_value"] for chunk in result["chunks"]}
        self.assertEqual(locators, {"交通安排", "餐饮建议"})
        combined = "\n".join(chunk["chunk_text"] for chunk in result["chunks"])
        self.assertNotIn("**", combined)
        self.assertNotIn("https://", combined)
        self.assertNotIn("dangerous", combined)
        self.assertTrue(all(chunk["page_number"] is None for chunk in result["chunks"]))

    def test_empty_markdown_is_rejected(self):
        with self.assertRaisesRegex(
            processor.DocumentProcessingError,
            "没有可用的文本",
        ):
            processor.process_document(b"# Empty\n", "empty.md")

    def test_document_id_is_stable_and_uses_original_bytes(self):
        content = ("稳定的文档标识测试内容。" * 10).encode("utf-8")
        first = processor.process_document(content, "first.txt")
        second = processor.process_document(content, "renamed.txt")

        self.assertEqual(first["document_id"], second["document_id"])

    def test_chunk_indexes_and_ids_are_unique(self):
        content = ("A" * 260).encode("utf-8")
        result = processor.process_document(
            content,
            "long.txt",
            target_size=100,
            overlap=20,
            min_size=20,
        )
        chunks = result["chunks"]

        self.assertEqual(
            [chunk["chunk_index"] for chunk in chunks],
            list(range(len(chunks))),
        )
        self.assertEqual(len({chunk["chunk_id"] for chunk in chunks}), len(chunks))
        self.assertEqual(chunks[0]["chunk_text"][-20:], chunks[1]["chunk_text"][:20])

    def test_path_like_upload_name_is_reduced_to_basename(self):
        content = ("安全文件名测试。" * 10).encode("utf-8")
        result = processor.process_document(content, "../../private/notes.txt")
        self.assertEqual(result["source_file"], "notes.txt")

    def test_document_size_limit_is_enforced_in_processor(self):
        oversized = b"x" * (processor.MAX_DOCUMENT_SIZE_BYTES + 1)
        with self.assertRaisesRegex(processor.DocumentProcessingError, "10MB"):
            processor.process_document(oversized, "large.txt")


class VectorStoreCompatibilityTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root_patch = patch.object(
            vector_store,
            "KNOWLEDGE_BASE_ROOT",
            Path(self.temporary_directory.name),
        )
        self.root_patch.start()

    def tearDown(self):
        self.root_patch.stop()
        self.temporary_directory.cleanup()

    def _write_legacy_index(self) -> None:
        index_path = vector_store._index_path(
            TEST_VISITOR_ID,
            create_directory=True,
        )
        np.savez_compressed(
            index_path,
            embeddings=np.asarray([[1.0, 0.0]], dtype=np.float32),
            chunk_texts=np.asarray(["旧 PDF 内容"], dtype=np.str_),
            document_ids=np.asarray(["a" * 64], dtype=np.str_),
            chunk_ids=np.asarray(["b" * 64], dtype=np.str_),
            source_files=np.asarray(["legacy.pdf"], dtype=np.str_),
            page_numbers=np.asarray([3], dtype=np.int32),
            chunk_indexes=np.asarray([0], dtype=np.int32),
        )

    def test_legacy_pdf_index_gets_compatible_locator_defaults(self):
        self._write_legacy_index()
        results = vector_store.search(
            TEST_VISITOR_ID,
            np.asarray([1.0, 0.0], dtype=np.float32),
            top_k=1,
        )

        self.assertEqual(results[0]["source_type"], "pdf")
        self.assertEqual(results[0]["page_number"], 3)
        self.assertEqual(results[0]["locator_type"], "page")
        self.assertEqual(results[0]["locator_value"], "3")

    def test_new_non_pdf_metadata_round_trips(self):
        result = processor.process_document(
            ("知识库 TXT 测试内容。" * 10).encode("utf-8"),
            "notes.txt",
        )
        vector_store.add_chunks(
            TEST_VISITOR_ID,
            result["chunks"],
            np.asarray([[1.0, 0.0]] * len(result["chunks"]), dtype=np.float32),
        )
        search_result = vector_store.search(
            TEST_VISITOR_ID,
            np.asarray([1.0, 0.0], dtype=np.float32),
            top_k=1,
        )[0]

        self.assertEqual(search_result["source_type"], "txt")
        self.assertIsNone(search_result["page_number"])
        self.assertEqual(search_result["locator_type"], "chunk")
        self.assertEqual(search_result["locator_value"], "1")

    def test_all_supported_document_types_coexist_and_are_searchable(self):
        documents = [
            processor.process_document(
                make_pdf_bytes(
                    "PDF knowledge about museum opening arrangements. " * 3
                ),
                "museum.pdf",
            ),
            processor.process_document(make_docx_bytes(), "travel.docx"),
            processor.process_document(
                ("这是 TXT 知识库内容，用于检索城市交通建议。" * 4).encode("utf-8"),
                "transport.txt",
            ),
            processor.process_document(
                (
                    "# 餐饮建议\n\n"
                    + "Markdown 资料建议提前确认餐厅的营业时间和预订规则。" * 4
                ).encode("utf-8"),
                "food.md",
            ),
        ]

        for dimension, document in enumerate(documents):
            embedding = np.zeros((len(document["chunks"]), 4), dtype=np.float32)
            embedding[:, dimension] = 1.0
            vector_store.add_chunks(
                TEST_VISITOR_ID,
                document["chunks"],
                embedding,
            )

        for dimension, expected_type in enumerate(("pdf", "docx", "txt", "md")):
            query = np.zeros(4, dtype=np.float32)
            query[dimension] = 1.0
            result = vector_store.search(TEST_VISITOR_ID, query, top_k=1)[0]
            self.assertEqual(result["source_type"], expected_type)
            self.assertIsNotNone(result["locator_type"])
            self.assertTrue(result["locator_value"])

        for document in documents:
            self.assertTrue(
                vector_store.has_document(
                    TEST_VISITOR_ID,
                    document["document_id"],
                )
            )


if __name__ == "__main__":
    unittest.main()
